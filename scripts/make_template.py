import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

# ─── helpers ─────────────────────────────────────────────
def fill(hex_color):
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type='solid')

def thin_border():
    s = Side(style='thin', color='BBBBBB')
    return Border(left=s, right=s, top=s, bottom=s)

def center(wrap=False):
    return Alignment(horizontal='center', vertical='center', wrap_text=wrap)

def left_align(wrap=False):
    return Alignment(horizontal='left', vertical='center', wrap_text=wrap)

# ─── colors ──────────────────────────────────────────────
C_NAVY    = fill('1F497D')
C_BLUE    = fill('2E75B6')
C_BLUE2   = fill('4472C4')
C_GREEN   = fill('375623')
C_ORANGE  = fill('C55A11')
C_AUTO    = fill('D9E1F2')
C_SAMPLE  = fill('FFF2CC')
C_GREY    = fill('F2F2F2')
C_WHITE   = fill('FFFFFF')

def wbold(size=11):  return Font(color='FFFFFF', bold=True,  name='TH SarabunPSK', size=size)
def wnorm(size=10):  return Font(color='FFFFFF', bold=False, name='TH SarabunPSK', size=size)
def bbold(size=10):  return Font(color='000000', bold=True,  name='TH SarabunPSK', size=size)
def bnorm(size=10):  return Font(color='000000', bold=False, name='TH SarabunPSK', size=size)
def grey_f(size=10): return Font(color='595959', bold=False, name='TH SarabunPSK', size=size)

# ─────────────────────────────────────────────────────────
# WORKBOOK
# ─────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Template'
ws.sheet_view.showGridLines = True
ws.freeze_panes = 'B5'

# ── Row 1 : Title ─────────────────────────────────────────
ws.row_dimensions[1].height = 28
ws.merge_cells('A1:AE1')
c = ws['A1']
c.value = 'แม่แบบนำเข้าข้อมูลโครงงาน — CSTU SPACE  |  เทอม 2  ปีการศึกษา 2568'
c.fill  = C_NAVY
c.font  = Font(color='FFFFFF', bold=True, name='TH SarabunPSK', size=14)
c.alignment = center()

# ── Row 2 : Section headers ───────────────────────────────
ws.row_dimensions[2].height = 22
SECTIONS = [
    ('A2',  'A2',  'เลขที่',              C_AUTO,   Font(color='595959', bold=True, name='TH SarabunPSK', size=10)),
    ('B2',  'G2',  'ข้อมูลโครงงาน',       C_NAVY,   wbold()),
    ('H2',  'M2',  'สมาชิกโครงงาน 1',     C_BLUE,   wbold()),
    ('N2',  'S2',  'สมาชิกโครงงาน 2',     C_BLUE2,  wbold()),
    ('T2',  'T2',  'สมาชิก',              C_AUTO,   Font(color='595959', bold=True, name='TH SarabunPSK', size=10)),
    ('U2',  'W2',  'อาจารย์ที่ปรึกษา',    C_GREEN,  wbold()),
    ('X2',  'AA2', 'รหัสอาจารย์/กรรมการ', C_GREEN,  wbold()),
    ('AB2', 'AE2', 'ตารางสอบ',            C_ORANGE, wbold()),
]
for start, end, label, bg, font in SECTIONS:
    if start != end:
        ws.merge_cells(f'{start}:{end}')
    c = ws[start]
    c.value = label; c.fill = bg; c.font = font; c.alignment = center()

# ── Row 3 : Column headers ────────────────────────────────
ws.row_dimensions[3].height = 44
HEADERS = [
    # col, label,                          bg,       font,    width, kind
    ('A',  'เลขที่\n(ระบบกำหนด)',          C_AUTO,   grey_f(),  9,   'auto'),
    ('B',  'รหัสโครงงาน\n(ProjCode) *',    C_NAVY,   wbold(), 22,   'req'),
    ('C',  'ชื่อย่อ ทปษ\n(= AdvId) *',     C_NAVY,   wbold(), 10,   'req'),
    ('D',  'ประเภท\nโครงงาน *',            C_NAVY,   wbold(), 10,   'req'),
    ('E',  'รหัสวิชา *',                   C_NAVY,   wbold(), 10,   'req'),
    ('F',  'ชื่อโครงงาน\nภาษาไทย *',       C_NAVY,   wbold(), 36,   'req'),
    ('G',  'ชื่อโครงงาน\nภาษาอังกฤษ',      C_NAVY,   wnorm(), 40,   'opt'),
    # member 1
    ('H',  'คำนำหน้า *',                   C_BLUE,   wbold(), 10,   'req'),
    ('I',  'ชื่อ-นามสกุล *',               C_BLUE,   wbold(), 24,   'req'),
    ('J',  'รหัสนักศึกษา *',               C_BLUE,   wbold(), 14,   'req'),
    ('K',  'อีเมล *',                       C_BLUE,   wbold(), 32,   'req'),
    ('L',  'เบอร์โทร',                      C_BLUE,   wnorm(), 13,   'opt'),
    ('M',  'ประเภท\nโครงการ *',             C_BLUE,   wbold(), 12,   'req'),
    # member 2
    ('N',  'คำนำหน้า',                     C_BLUE2,  wnorm(), 10,   'opt'),
    ('O',  'ชื่อ-นามสกุล',                 C_BLUE2,  wnorm(), 24,   'opt'),
    ('P',  'รหัสนักศึกษา',                 C_BLUE2,  wnorm(), 14,   'opt'),
    ('Q',  'อีเมล',                         C_BLUE2,  wnorm(), 32,   'opt'),
    ('R',  'เบอร์โทร',                      C_BLUE2,  wnorm(), 13,   'opt'),
    ('S',  'ประเภท\nโครงการ',               C_BLUE2,  wnorm(), 12,   'opt'),
    # auto
    ('T',  'จำนวน\nสมาชิก\n(ระบบกำหนด)',   C_AUTO,   grey_f(),  9,   'auto'),
    # advisors
    ('U',  'ชื่อ ทปษ\n(ระบบเติมให้)',       C_AUTO,   grey_f(), 24,   'auto'),
    ('V',  'ที่ปรึกษาร่วม\nภายใน\n(ระบบเติมให้)', C_AUTO, grey_f(), 24, 'auto'),
    ('W',  'ที่ปรึกษาร่วม\nภายนอก\n(กรอกเอง)',    C_GREEN, wnorm(), 24, 'opt'),
    # codes
    ('X',  'AdvId *',                       C_GREEN,  wbold(), 10,   'req'),
    ('Y',  'Comm1 *',                       C_GREEN,  wbold(), 10,   'req'),
    ('Z',  'Comm2',                         C_GREEN,  wnorm(), 10,   'opt'),
    ('AA', 'Comm3',                         C_GREEN,  wnorm(), 10,   'opt'),
    # schedule
    ('AB', 'วันสอบ *\n(YYYY-MM-DD)',         C_ORANGE, wbold(), 16,   'req'),
    ('AC', 'เวลาเริ่ม *\n(HH:MM)',           C_ORANGE, wbold(), 11,   'req'),
    ('AD', 'เวลาจบ *\n(HH:MM)',              C_ORANGE, wbold(), 11,   'req'),
    ('AE', 'ห้องสอบ',                        C_ORANGE, wnorm(), 14,   'opt'),
]
for col, label, bg, font, width, kind in HEADERS:
    c = ws[f'{col}3']
    c.value = label
    c.fill  = bg
    c.font  = font
    c.alignment = center(wrap=True)
    c.border = thin_border()
    ws.column_dimensions[col].width = width

# ── Row 4 : Legend ────────────────────────────────────────
ws.row_dimensions[4].height = 18
ws.merge_cells('A4:G4')
c = ws['A4']
c.value = '* = จำเป็นต้องกรอก    แถวสีเหลือง = ตัวอย่างข้อมูล    ช่องสีฟ้าอ่อน = ระบบกำหนดอัตโนมัติ ไม่ต้องกรอก'
c.font  = Font(color='7F6000', italic=True, name='TH SarabunPSK', size=9)
c.alignment = left_align()

# ── Row 5 : Sample data ───────────────────────────────────
ws.row_dimensions[5].height = 18
SAMPLE = [
    '', '68-2_01_kdc-s1', 'kdc', 's', 'CS303',
    'ระบบแจ้งและติดตามปัญหาภายในมหาวิทยาลัย',
    'UNIVERSITY ISSUE REPORTING AND TRACKING SYSTEM',
    'นาย', 'ณธรกริช ทองธรรมชาติ', '6609650293',
    'nathornkrich.thon@dome.tu.ac.th', '0917087755', 's',
    '', '', '', '', '', '',
    '1', 'ผศ. ดร.กษิดิศ ชาญเชี่ยว', '', '',
    'kdc', 'ssr', 'wlr', '',
    '2025-11-20', '13:00', '14:00', 'ห้อง CS 123',
]
for (col, *_), val in zip(HEADERS, SAMPLE):
    c = ws[f'{col}5']
    c.value = val
    c.fill  = C_SAMPLE
    c.font  = Font(color='7F6000', italic=True, name='TH SarabunPSK', size=10)
    c.alignment = left_align()
    c.border = thin_border()

# ── Rows 6–85 : Empty data rows ───────────────────────────
for row in range(6, 86):
    ws.row_dimensions[row].height = 18
    for col, label, bg, font, width, kind in HEADERS:
        c = ws[f'{col}{row}']
        c.fill   = C_GREY if kind == 'auto' else C_WHITE
        c.font   = grey_f() if kind == 'auto' else bnorm()
        c.alignment = left_align()
        c.border = thin_border()

# ── Data Validation ───────────────────────────────────────
def add_dv(ws, type_, formula, cells, prompt='', error=''):
    dv = DataValidation(type=type_, formula1=formula, allow_blank=True)
    if prompt: dv.prompt = prompt
    if error:  dv.error  = error
    ws.add_data_validation(dv)
    for cell_range in cells:
        dv.add(cell_range)

add_dv(ws, 'list', '"s,r,m"',         ['D6:D85'],
       prompt='s=ภาคพิเศษ, r=ปกติ, m=ผสม',
       error='ต้องเป็น s, r, หรือ m เท่านั้น')
add_dv(ws, 'list', '"CS303,CS403"',    ['E6:E85'])
add_dv(ws, 'list', '"นาย,นางสาว,นาง"',['H6:H85', 'N6:N85'])
add_dv(ws, 'list', '"s,r"',            ['M6:M85', 'S6:S85'],
       prompt='s=ภาคพิเศษ  r=ปกติ')

# ─────────────────────────────────────────────────────────
# SHEET 2 : คำแนะนำ
# ─────────────────────────────────────────────────────────
ws2 = wb.create_sheet('คำแนะนำ')
ws2.column_dimensions['A'].width = 24
ws2.column_dimensions['B'].width = 46
ws2.column_dimensions['C'].width = 30

def write_section(ws, row, title, rows_data, bg_title=C_NAVY):
    ws.row_dimensions[row].height = 22
    ws.merge_cells(f'A{row}:C{row}')
    c = ws[f'A{row}']
    c.value = title; c.fill = bg_title; c.font = wbold(); c.alignment = left_align()
    row += 1
    for a, b, cv in rows_data:
        ws.row_dimensions[row].height = 18
        for col_letter, val, fnt in [('A', a, bbold()), ('B', b, bnorm()), ('C', cv, grey_f())]:
            cell = ws[f'{col_letter}{row}']
            cell.value = val; cell.font = fnt
            cell.fill = fill('F9F9F9'); cell.border = thin_border()
            cell.alignment = left_align()
        row += 1
    return row + 1

# title
ws2.row_dimensions[1].height = 28
ws2.merge_cells('A1:C1')
c = ws2['A1']
c.value = 'คำแนะนำการกรอกข้อมูล — Import Template'
c.fill = C_NAVY
c.font = Font(color='FFFFFF', bold=True, name='TH SarabunPSK', size=13)
c.alignment = center()

r = 3
r = write_section(ws2, r, 'ประเภทโครงงาน (Project Type)', [
    ('s', 'ภาคพิเศษ (Special)',  'โครงการปริญญาตรีภาคพิเศษ'),
    ('r', 'ปกติ (Regular)',       'โครงการปกติ'),
    ('m', 'ผสม (Mixed)',          'สมาชิก 2 คน ต่างประเภทกัน'),
])
r = write_section(ws2, r, 'รหัสวิชา', [
    ('CS303', 'โครงงานพิเศษ 1', 'เทอม 2 ปีการศึกษา 2568'),
    ('CS403', 'โครงงานพิเศษ 2', 'เทอม 2 ปีการศึกษา 2568'),
])
r = write_section(ws2, r, 'รหัสอาจารย์ทั้งหมดในระบบ (ใช้ใน AdvId, Comm1-3)', [
    ('kdc',  'กษิดิศ ชาญเชี่ยว',          'ผศ. ดร.'),
    ('skn',  'สิริกันยา นิลพานิช',         'อ.'),
    ('ssr',  'ทรงศักดิ์ รองวิริยะพานิช',  'ผศ. ดร.'),
    ('wlr',  'วิลาวรรณ รักผกาวงศ์',       'ผศ.'),
    ('lpp',  'ลัมพาพรรณ พันธุ์ซูจิตร์',   'ผศ. ดร.'),
    ('nrc',  'นวฤกษ์ ชลารักษ์',           'อ.'),
    ('ojs',  'อรจิรา สิทธิศักดิ์',         'ผศ.'),
    ('pkl',  'ปกรณ์ ลี้สุทธิพรชัย',        'อ.'),
    ('pkp',  'ภัคพร เสาร์ฝืน',             'อ.'),
    ('ppr',  'ประภาพร รัตนธำรง',           'อ.'),
    ('pps',  'ปกป้อง ส่องเมือง',           'ผศ.'),
    ('scw',  'เสาวลักษณ์ วรรธนาภา',       'รศ. ดร.'),
    ('snk',  'ศาตนาฏ กิจศิรานุวัตร',       'อ.'),
    ('tnt',  'ธนาธร ทะนานทอง',             'อ.'),
    ('tpb',  'ฐาปนา บุญชู',                'ผศ.'),
    ('wdp',  'วนิดา พฤทธิวิทยา',           'ผศ. ดร.'),
    ('wjr',  'วิรัตน์ จารีวงศ์ไพบูลย์',    'อ.'),
    ('ddp',  'เด่นดวง ประดับสุวรรณ',       'รศ. ดร.'),
    ('nth',  'placeholder (ยังไม่ระบุ)',    'อ.'),
])
r = write_section(ws2, r, 'รูปแบบวันที่และเวลา', [
    ('วันสอบ',   'YYYY-MM-DD',  'เช่น 2025-11-18'),
    ('เวลาเริ่ม', 'HH:MM',      'เช่น 08:00'),
    ('เวลาจบ',   'HH:MM',       'เช่น 09:00'),
])
r = write_section(ws2, r, 'คอลัมน์ที่ระบบกำหนดอัตโนมัติ (ไม่ต้องกรอก — ช่องสีฟ้าอ่อน)', [
    ('A : เลขที่',         'ระบบนับให้อัตโนมัติ',        '-'),
    ('T : จำนวนสมาชิก',    'ระบบนับจากสมาชิก 1-2',      '-'),
    ('U : ชื่อ ทปษ',       'ระบบดึงจาก AdvId',           '-'),
    ('V : ที่ปรึกษาร่วมใน', 'ระบบดึงจาก Comm ที่ตรงกัน', '-'),
])
r = write_section(ws2, r, 'การตรวจสอบข้อมูลซ้ำ (ก่อน import)', [
    ('ProjCode ซ้ำ',          'แจ้ง warning — เลือก skip หรือ overwrite', ''),
    ('รหัสนักศึกษาซ้ำ',       'upsert — อัปเดตข้อมูลล่าสุด',             ''),
    ('รหัสอาจารย์ไม่มีในระบบ', 'Error — ห้าม import จนกว่าจะเพิ่ม',       'ดูรหัสที่ถูกต้องด้านบน'),
    ('นักศึกษาอยู่กลุ่มอื่นแล้ว','Warning — coordinator ตัดสิน',           ''),
])

# ─────────────────────────────────────────────────────────
# SAVE
# ─────────────────────────────────────────────────────────
path = r'C:/Users/U S E R/CSTU_SPACE_PHASE1/public/templates/project_import_template.xlsx'
wb.save(path)
print('Saved:', path)
